"""
TwinCore X - Utility functions for OCR, preprocessing, similarity analysis, and reporting.
Public-ready version: supports uploaded PDF files.
"""

import os
import re
import tempfile
import shutil
from typing import Dict, List, Tuple, Optional, Any
import numpy as np
import pandas as pd
from sklearn.metrics.pairwise import cosine_similarity
from sentence_transformers import SentenceTransformer
from paddleocr import PaddleOCR
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import streamlit as st


# ---------------------------------------------------------------------------
# OCR Engine (lazy loaded)
# ---------------------------------------------------------------------------
_ocr_engine: Optional[PaddleOCR] = None
_embedding_model: Optional[SentenceTransformer] = None


def get_ocr_engine() -> PaddleOCR:
    """Lazy-load PaddleOCR engine (English)."""
    global _ocr_engine
    if _ocr_engine is None:
        _ocr_engine = PaddleOCR(
            use_angle_cls=True,
            lang="en",
            show_log=False,
            use_gpu=False,
        )
    return _ocr_engine


def get_embedding_model() -> SentenceTransformer:
    """Lazy-load sentence-transformers model."""
    global _embedding_model
    if _embedding_model is None:
        _embedding_model = SentenceTransformer("all-MiniLM-L6-v2")
    return _embedding_model


# ---------------------------------------------------------------------------
# Text Extraction
# ---------------------------------------------------------------------------
def extract_text_from_pdf(pdf_path: str) -> str:
    """
    Extract text from a PDF using PaddleOCR.
    Handles multi-page PDFs. Returns concatenated text.
    """
    try:
        ocr = get_ocr_engine()
        result = ocr.ocr(pdf_path, cls=True)

        texts = []
        if result is None:
            return ""

        for page_result in result:
            if page_result is None:
                continue
            for line in page_result:
                if line and len(line) >= 2:
                    text = line[1][0]
                    if text and isinstance(text, str):
                        texts.append(text.strip())

        return " ".join(texts)
    except Exception as e:
        st.warning(f"OCR failed for {os.path.basename(pdf_path)}: {str(e)}")
        return ""


def extract_texts_from_uploaded_files(uploaded_files: List[Any]) -> Dict[str, str]:
    """
    Save uploaded PDF files to a temporary directory, run OCR, then clean up.
    Returns dict: {filename: extracted_text}
    """
    if not uploaded_files:
        raise ValueError("No files uploaded.")

    texts = {}
    temp_dir = tempfile.mkdtemp(prefix="twincorex_")

    try:
        progress = st.progress(0, text="Extracting text from PDFs...")
        total = len(uploaded_files)

        for idx, uploaded_file in enumerate(uploaded_files):
            filename = uploaded_file.name
            # Save to temp file
            temp_path = os.path.join(temp_dir, filename)
            with open(temp_path, "wb") as f:
                f.write(uploaded_file.getbuffer())

            text = extract_text_from_pdf(temp_path)
            texts[filename] = text
            progress.progress(
                (idx + 1) / total,
                text=f"Processed {idx + 1}/{total}: {filename}"
            )

        progress.empty()
    finally:
        # Always clean up temp files
        shutil.rmtree(temp_dir, ignore_errors=True)

    return texts


# Keep old function for local use if needed
def extract_texts_from_folder(folder_path: str) -> Dict[str, str]:
    """Legacy: Read all PDF files from a local folder."""
    import glob
    if not os.path.isdir(folder_path):
        raise ValueError(f"Invalid folder path: {folder_path}")

    pdf_files = sorted(glob.glob(os.path.join(folder_path, "*.pdf")))
    if not pdf_files:
        raise ValueError(f"No PDF files found in: {folder_path}")

    texts = {}
    progress = st.progress(0, text="Extracting text from PDFs...")
    total = len(pdf_files)

    for idx, pdf_path in enumerate(pdf_files):
        filename = os.path.basename(pdf_path)
        text = extract_text_from_pdf(pdf_path)
        texts[filename] = text
        progress.progress((idx + 1) / total, text=f"Processed {idx + 1}/{total}: {filename}")

    progress.empty()
    return texts


# ---------------------------------------------------------------------------
# Preprocessing
# ---------------------------------------------------------------------------
def clean_text(text: str) -> str:
    """Remove extra whitespace, control characters and normalize."""
    if not text:
        return ""
    text = re.sub(r"[^\x20-\x7E\n\t]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def preprocess_texts(raw_texts: Dict[str, str]) -> Dict[str, str]:
    """Apply cleaning to all extracted texts."""
    return {fname: clean_text(txt) for fname, txt in raw_texts.items()}


# ---------------------------------------------------------------------------
# Similarity Analysis
# ---------------------------------------------------------------------------
def compute_embeddings(texts: Dict[str, str]) -> Tuple[List[str], np.ndarray]:
    """Compute sentence embeddings for all student answers."""
    model = get_embedding_model()
    filenames = list(texts.keys())
    corpus = [texts[f] if texts[f] else " " for f in filenames]
    embeddings = model.encode(corpus, show_progress_bar=False, convert_to_numpy=True)
    return filenames, embeddings


def compute_similarity_matrix(
    filenames: List[str], embeddings: np.ndarray
) -> pd.DataFrame:
    """Compute pairwise cosine similarity matrix as DataFrame."""
    sim_matrix = cosine_similarity(embeddings)
    df = pd.DataFrame(sim_matrix, index=filenames, columns=filenames)
    return df


def get_suspicious_pairs(
    sim_df: pd.DataFrame, threshold: float = 0.75
) -> List[Tuple[str, str, float]]:
    """Return list of (student_a, student_b, similarity) where sim >= threshold."""
    pairs = []
    files = sim_df.index.tolist()
    for i in range(len(files)):
        for j in range(i + 1, len(files)):
            score = float(sim_df.iloc[i, j])
            if score >= threshold:
                pairs.append((files[i], files[j], score))
    pairs.sort(key=lambda x: x[2], reverse=True)
    return pairs


# ---------------------------------------------------------------------------
# Report Generation
# ---------------------------------------------------------------------------
class EvidencePDF(FPDF):
    """Custom PDF for evidence report."""

    def header(self):
        self.set_font("Helvetica", "B", 16)
        self.cell(0, 10, "TwinCore X - Evidence Report", ln=True, align="C")
        self.set_font("Helvetica", "", 10)
        self.cell(0, 8, "Student Assignment Similarity Analysis", ln=True, align="C")
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")


def generate_pdf_report(
    total_pdfs: int,
    suspicious_pairs: List[Tuple[str, str, float]],
    texts: Dict[str, str],
    threshold: float = 0.75,
) -> bytes:
    """Generate evidence PDF report and return as bytes."""
    pdf = EvidencePDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "1. Summary", ln=True)
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 7, f"Total PDFs analyzed: {total_pdfs}", ln=True)
    pdf.cell(0, 7, f"Similarity threshold: {threshold * 100:.0f}%", ln=True)
    pdf.cell(0, 7, f"Suspicious pairs found: {len(suspicious_pairs)}", ln=True)
    pdf.ln(5)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "2. Flagged Suspicious Pairs", ln=True)
    pdf.set_font("Helvetica", "", 10)

    if not suspicious_pairs:
        pdf.cell(0, 7, "No pairs exceeded the similarity threshold.", ln=True)
    else:
        pdf.set_fill_color(200, 220, 255)
        pdf.cell(70, 8, "Student A", border=1, fill=True)
        pdf.cell(70, 8, "Student B", border=1, fill=True)
        pdf.cell(40, 8, "Similarity %", border=1, fill=True, ln=True)

        for a, b, score in suspicious_pairs:
            pdf.cell(70, 7, a[:40], border=1)
            pdf.cell(70, 7, b[:40], border=1)
            pdf.cell(40, 7, f"{score * 100:.1f}%", border=1, ln=True)

    pdf.ln(8)

    if suspicious_pairs:
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "3. Text Excerpts (first 400 characters)", ln=True)
        pdf.set_font("Helvetica", "", 9)

        for idx, (a, b, score) in enumerate(suspicious_pairs, 1):
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(0, 7, f"Pair {idx}: {a}  vs  {b}  ({score * 100:.1f}%)", ln=True)
            pdf.set_font("Helvetica", "", 8)

            text_a = texts.get(a, "")[:400]
            text_b = texts.get(b, "")[:400]

            pdf.multi_cell(0, 5, f"[{a}]: {text_a}...")
            pdf.ln(1)
            pdf.multi_cell(0, 5, f"[{b}]: {text_b}...")
            pdf.ln(2)

    pdf.ln(5)
    pdf.set_font("Helvetica", "I", 8)
    pdf.multi_cell(
        0,
        5,
        "Disclaimer: This report is generated by TwinCore X using semantic similarity "
        "analysis (sentence embeddings + cosine similarity). High similarity does not "
        "automatically prove plagiarism; it is intended as a screening aid for teachers. "
        "Manual review is always recommended.",
    )

    return bytes(pdf.output())


def generate_excel_report(
    suspicious_pairs: List[Tuple[str, str, float]], threshold: float = 0.75
) -> bytes:
    """Generate Excel file of suspicious pairs and return as bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Suspicious Pairs"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["#", "Student A", "Student B", "Similarity (%)", "Flagged"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for idx, (a, b, score) in enumerate(suspicious_pairs, 1):
        ws.cell(row=idx + 1, column=1, value=idx).border = thin_border
        ws.cell(row=idx + 1, column=2, value=a).border = thin_border
        ws.cell(row=idx + 1, column=3, value=b).border = thin_border
        cell = ws.cell(row=idx + 1, column=4, value=round(score * 100, 2))
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")
        flag_cell = ws.cell(row=idx + 1, column=5, value="YES" if score >= threshold else "NO")
        flag_cell.border = thin_border
        flag_cell.alignment = Alignment(horizontal="center")
        if score >= 0.9:
            flag_cell.fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
        elif score >= 0.8:
            flag_cell.fill = PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid")

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 12

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "TwinCore X - Similarity Analysis Summary"
    ws2["A1"].font = Font(bold=True, size=14)
    ws2["A3"] = "Threshold used:"
    ws2["B3"] = f"{threshold * 100:.0f}%"
    ws2["A4"] = "Total suspicious pairs:"
    ws2["B4"] = len(suspicious_pairs)
    ws2["A6"] = "Generated by TwinCore X"
    ws2["A6"].font = Font(italic=True, size=9)

    from io import BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
